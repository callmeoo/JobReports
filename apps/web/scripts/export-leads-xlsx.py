#!/usr/bin/env python3
"""Export 线索验证.xlsx → JSON for Prisma buyer import."""

import json
from pathlib import Path

import openpyxl

SRC = Path("/Users/yuanlijuan/Desktop/TMP/线索验证.xlsx")
OUT = Path("/Users/yuanlijuan/Documents/Cursor/JiJi/apps/web/tmp-leads-import.json")

PRIORITY_TIER = {
    "A+": "KEY",
    "A": "KEY",
    "B": "FOLLOW",
    "C": "WATCH",
    "D": "LOW",
}


def cell(row, idx, name):
    i = idx.get(name)
    if i is None:
        return None
    v = row[i]
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["线索验证"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    leads = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        company = cell(row, idx, "company_name")
        seller_id = cell(row, idx, "seller_id")
        if not company:
            continue

        phone = cell(row, idx, "phone_number")
        wa = cell(row, idx, "whatsapp") or phone
        location = cell(row, idx, "location") or ""
        city = location.split(",")[0].strip() if location else None
        sub = cell(row, idx, "lagos_sub_area")
        if sub and city:
            city = f"{city} / {sub}"

        priority = cell(row, idx, "priority_level") or ""
        score = cell(row, idx, "score")
        biz = cell(row, idx, "business_type_judgement")
        wholesale = cell(row, idx, "wholesale_potential")
        brands = cell(row, idx, "main_brands")
        ads = cell(row, idx, "cars_ads_count")
        url = cell(row, idx, "seller_url")
        reasons = cell(row, idx, "b2b_match_reasons")
        existing_notes = cell(row, idx, "notes")

        note_parts = [
            f"priority={priority}",
            f"score={score}" if score else None,
            f"type={biz}" if biz else None,
            f"wholesale={wholesale}" if wholesale else None,
            f"ads={ads}" if ads else None,
            f"brands={brands}" if brands else None,
            f"url={url}" if url else None,
            f"match={reasons}" if reasons else None,
            existing_notes,
        ]
        notes = "\n".join(p for p in note_parts if p)

        tags = []
        if priority:
            tags.append(priority)
        if biz:
            tags.append(biz)
        if wholesale:
            tags.append(f"批发:{wholesale}")

        leads.append(
            {
                "companyName": company,
                "whatsapp": wa,
                "phone": phone,
                "countryCode": "NG",
                "city": city,
                "source": "JIJI",
                "sourceRef": f"jiji:{seller_id}" if seller_id else f"jiji-name:{company}",
                "stage": "NEW",
                "tier": PRIORITY_TIER.get(priority, "LOW"),
                "notes": notes,
                "tags": tags,
            }
        )

    OUT.write_text(json.dumps(leads, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"exported {len(leads)} → {OUT}")
    with_phone = sum(1 for x in leads if x["whatsapp"])
    print(f"with phone/whatsapp: {with_phone}")
    print(f"KEY tier: {sum(1 for x in leads if x['tier']=='KEY')}")


if __name__ == "__main__":
    main()
